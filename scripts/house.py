"""Detached house prices per map area.

England and Wales: median price of standard-market (category A) detached sales in the last
36 months of Land Registry Price Paid Data, each sale placed by its postcode (Code-Point Open)
in a map area. Areas with too few sales take their whole council's median.
Scotland: UK HPI detached average price for the council, mean of the last 12 months.
Northern Ireland: no detached price by council is published, so NI's detached standardised
price is scaled by each council's all-property price relative to NI's (NISRA NI HPI).
Both are scaled by the median ratio of sales median to HPI detached price across English and
Welsh councils, so the nations are comparable.
"""
import csv, glob, pickle, statistics
from collections import defaultdict
from datetime import date
import numpy as np
import openpyxl
from shapely import STRtree, from_wkb, points

MIN_SALES = 30
MONTHS = 36

areas = list(csv.DictReader(open('met/uk_areas.csv')))
parent_of = {a['id']: a['id'].split('-')[0] for a in areas}
shapes = pickle.load(open('grid/area_shapes.pkl', 'rb'))
ids = [i for i, _ in shapes]
tree = STRtree([from_wkb(w) for _, w in shapes])

pc = {}
for f in glob.glob('house/codepo/Data/CSV/*.csv'):
    for r in csv.reader(open(f)):
        if r[1] != '90':  # quality 90: no coordinates
            pc[r[0].replace(' ', '')] = (float(r[2]), float(r[3]))

sales = []
for f in sorted(glob.glob('house/pp-*.csv')):
    for r in csv.reader(open(f)):
        if r[4] == 'D' and r[14] == 'A' and r[15] != 'D':
            sales.append((r[2][:10], int(r[1]), r[3].replace(' ', '')))
latest = max(s[0] for s in sales)
y, m = int(latest[:4]), int(latest[5:7])
# window: the 36 whole months ending with the latest month in the data
first = date(y - 3 + (m == 12), m % 12 + 1, 1).isoformat()
sales = [s for s in sales if s[0] >= first]

xy = np.array([pc.get(s[2], (np.nan, np.nan)) for s in sales])
ok = ~np.isnan(xy[:, 0])
hit_pt, hit_area = tree.query(points(xy[ok]), predicate='intersects')
area_of = np.full(len(sales), -1)
area_of[np.nonzero(ok)[0][hit_pt]] = hit_area
by_area, by_council = defaultdict(list), defaultdict(list)
for s, a in zip(sales, area_of):
    if a >= 0:
        by_area[ids[a]].append(s[1])
        by_council[parent_of[ids[a]]].append(s[1])
print(f'{len(sales)} detached sales {first} to {latest}; placed {sum(1 for a in area_of if a >= 0)}; no postcode match {int((~ok).sum())}')

# Scotland: UK HPI detached average, last 12 months
hpi = defaultdict(list)
for r in csv.DictReader(open('house/hpi.csv')):
    d = r['Date'][6:] + '-' + r['Date'][3:5]
    if r['AreaCode'][:3] in ('S12', 'E06', 'E07', 'E08', 'E09', 'W06') and r['DetachedPrice'] and d > f'{int(latest[:4]) - 1}-{latest[5:7]}':
        hpi[r['AreaCode']].append(float(r['DetachedPrice']))
# The HPI figure is a mix-adjusted average for the last year; put it on the same footing as the
# sales medians using the typical ratio between the two across English and Welsh councils.
calib = statistics.median(statistics.median(by_council[c]) / statistics.mean(v) for c, v in hpi.items()
                          if c[0] in 'EW' and len(by_council.get(c, [])) >= MIN_SALES)

# Northern Ireland: scale NI detached standardised price by each LGD's all-property price (latest 4 quarters)
wb = openpyxl.load_workbook('house/nihpi.xlsx', read_only=True, data_only=True)
det = [r[4] for r in wb['Table 9a'].iter_rows(min_row=4, values_only=True) if isinstance(r[4], (int, float))][-4:]
allp = [r[4] for r in wb['Table 9'].iter_rows(min_row=4, values_only=True) if isinstance(r[4], (int, float))][-4:]
ratio_ni = statistics.mean(det) / statistics.mean(allp)
t5 = list(wb['Table 5'].iter_rows(values_only=True))
head = next(i for i, r in enumerate(t5) if r and any(isinstance(c, str) and 'Belfast' in c for c in r))
cols = t5[head]
lgd_price = {}
for name in [c for c in cols if isinstance(c, str) and 'Price' in c]:
    vals = [r[cols.index(name)] for r in t5[head + 1:] if isinstance(r[cols.index(name)], (int, float))][-4:]
    lgd_price[name.replace(' Standardised Price', '').strip().lower()] = statistics.mean(vals)

out = []
for a in areas:
    nation, own, par = a['nation'], by_area.get(a['id'], []), by_council.get(parent_of[a['id']], [])
    price = sales_n = None
    basis = ''
    if nation in ('England', 'Wales'):
        if len(own) >= MIN_SALES: price, sales_n, basis = statistics.median(own), len(own), 'sales'
        elif len(par) >= MIN_SALES: price, sales_n, basis = statistics.median(par), len(par), 'council'
    elif nation == 'Scotland':
        v = hpi.get(parent_of[a['id']])
        if v: price, basis = statistics.mean(v) * calib, 'hpi'
    else:
        name = (a['parent'] or a['name']).lower()
        v = next((p for k, p in lgd_price.items() if k.startswith(name[:12])), None)
        if v: price, basis = v * ratio_ni * calib, 'estimate'
    out.append({'id': a['id'], 'price': round(price, -3) if price else '', 'sales': sales_n or '', 'basis': basis})

with open('met/uk_area_houses.csv', 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['id', 'price', 'sales', 'basis']); w.writeheader(); w.writerows(out)
from collections import Counter
print(Counter(r['basis'] for r in out), 'window', first, latest, 'NI ratio', round(ratio_ni, 3), 'HPI calibration', round(calib, 3), 'NI LGDs', len(lgd_price))
